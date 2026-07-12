"""Unit and integration tests for Advanced Analytics statistics modules."""

from datetime import date
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import ComplicationEvent, RiskFactor, Surgery
from app.statistics.analytics import build_advanced_analytics
from app.statistics.confidence_interval import wald_log_ratio_ci
from app.statistics.distribution import DistributionInput, suggest_distribution
from app.statistics.export import build_analytics_workbook
from app.statistics.fisher import fisher_exact_pvalue
from app.statistics.odds_ratio import compute_odds_ratio_vs_rest
from app.statistics.risk import build_risk_factor_stats
from tests.conftest import login


def add_case(db: Session, surgeon_id: int, creator_id: int, code: str, *, pcr=False, risks=()):
    surgery = Surgery(
        case_code=code,
        surgery_date=date(2026, 7, 1),
        eye="OD",
        surgeon_id=surgeon_id,
        created_by_id=creator_id,
    )
    surgery.risk_factors = [RiskFactor(code=risk) for risk in risks]
    surgery.complication = ComplicationEvent(
        occurred=pcr,
        complication_type="pcr" if pcr else None,
    )
    db.add(surgery)
    db.commit()


def test_odds_ratio_and_ci_with_haldane():
    result = compute_odds_ratio_vs_rest(2, 10, 0, 10)
    assert result.comparable is True
    assert result.used_haldane_correction is True
    assert result.odds_ratio is not None and result.odds_ratio > 1
    assert result.ci_low is not None and result.ci_high is not None
    assert result.ci_low < result.odds_ratio < result.ci_high
    assert result.p_value is not None
    assert result.relative_risk is not None
    assert result.absolute_risk_difference is not None


def test_odds_ratio_unavailable_without_comparator():
    result = compute_odds_ratio_vs_rest(1, 5, 0, 0)
    assert result.comparable is False
    assert result.odds_ratio is None
    assert result.p_value is None


def test_wald_ci_positive():
    low, high = wald_log_ratio_ci(2.0, se=0.3)
    assert low < 2.0 < high


def test_fisher_exact_known_table():
    # Classic 2x2 where association is strong.
    p = fisher_exact_pvalue(8, 2, 1, 9)
    assert 0 <= p <= 1
    assert p < 0.05


def test_distribution_sum_and_max_difference():
    surgeons = [
        DistributionInput(1, "A", 20, 1),
        DistributionInput(2, "B", 20, 4),
        DistributionInput(3, "C", 20, 2),
    ]
    rows = suggest_distribution(surgeons, total_cases=20, max_difference=3)
    counts = [row.suggested_cases for row in rows]
    assert sum(counts) == 20
    assert max(counts) - min(counts) <= 3
    # Lower smoothed risk should tend to receive more cases.
    by_name = {row.surgeon_name: row for row in rows}
    assert by_name["A"].suggested_cases >= by_name["B"].suggested_cases


def test_distribution_max_difference_zero_no_hang():
    surgeons = [
        DistributionInput(1, "A", 10, 1),
        DistributionInput(2, "B", 10, 3),
    ]
    rows = suggest_distribution(surgeons, total_cases=5, max_difference=0)
    counts = [row.suggested_cases for row in rows]
    assert sum(counts) == 5
    assert max(counts) - min(counts) == 1


def test_risk_factor_top_n_and_pareto():
    stats = build_risk_factor_stats(
        {"dense_cataract": 5, "small_pupil": 3, "other": 1},
        total_surgeries=10,
        catalog={"dense_cataract": "Catarata densa", "small_pupil": "Pupila pequeña", "other": "Otro"},
        top_n=2,
    )
    assert len(stats) == 2
    assert stats[0].code == "dense_cataract"
    assert stats[0].cumulative_pct <= stats[1].cumulative_pct
    assert stats[-1].cumulative_pct == 100.0


def test_analytics_bundle_and_export(db_session, seed_users):
    admin = seed_users["admin"]
    s1 = seed_users["surgeon"]
    s2 = seed_users["surgeon2"]
    for idx in range(10):
        add_case(db_session, s1.id, admin.id, f"A{idx}", pcr=idx == 0, risks=("small_pupil",) if idx < 3 else ())
    for idx in range(10):
        add_case(db_session, s2.id, admin.id, f"B{idx}", pcr=idx < 3, risks=("dense_cataract",) if idx < 5 else ())

    bundle = build_advanced_analytics(db_session, planned_cases=20, max_difference=3)
    assert bundle.summary.total_surgeries == 20
    assert bundle.summary.total_pcr == 4
    assert len(bundle.surgeon_rows) == 2
    assert sum(row.suggested_cases for row in bundle.distribution) == 20
    assert bundle.risk_factors

    payload = build_analytics_workbook(bundle)
    wb = load_workbook(BytesIO(payload))
    assert wb.sheetnames == ["Resumen", "OR", "Distribución", "Factores de riesgo", "Reintervenciones"]
    assert wb["OR"].max_row >= 3
    assert bundle.reintervention_types
    assert any(r.label == "Lente fijado a esclera" for r in bundle.reintervention_types)


def test_advanced_analytics_admin_only(client: TestClient, db_session, seed_users):
    add_case(db_session, seed_users["surgeon"].id, seed_users["admin"].id, "VIS1")

    login(client, "admin", "admin123")
    page = client.get("/admin/analytics")
    assert page.status_code == 200
    assert "Advanced Analytics" in page.text
    assert "Surgical Distribution Planner" in page.text
    assert "Reintervenciones por tipo general" in page.text
    assert "Lente fijado a esclera" in page.text

    export = client.get("/admin/analytics/export")
    assert export.status_code == 200
    assert "spreadsheetml" in export.headers["content-type"]

    login(client, "coord", "coord123")
    denied = client.get("/admin/analytics")
    assert denied.status_code == 403
