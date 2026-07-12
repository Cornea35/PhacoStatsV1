"""Excel export for Advanced Analytics (openpyxl)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.statistics.analytics import AdvancedAnalyticsBundle


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)


def build_analytics_workbook(bundle: AdvancedAnalyticsBundle) -> bytes:
    """Return an .xlsx payload with one sheet per analysis block."""
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Resumen"
    _write_header(ws_summary, ["Métrica", "Valor"])
    summary = bundle.summary
    rows = [
        ("Cirugías totales", summary.total_surgeries),
        ("RCP totales", summary.total_pcr),
        ("Tasa global RCP (%)", summary.global_pcr_rate_pct),
        ("Número de cirujanos", summary.surgeon_count),
        ("Odds Ratio promedio", summary.average_odds_ratio if summary.average_odds_ratio is not None else "N/A"),
        ("Factor de riesgo más frecuente", summary.top_risk_factor or "N/A"),
        ("Cirujano con menor RCP", summary.lowest_pcr_surgeon or "N/A"),
        ("Cirujano con mayor RCP", summary.highest_pcr_surgeon or "N/A"),
        ("Cirugías planificadas (input)", bundle.planned_cases),
        ("Diferencia máxima (input)", bundle.max_difference),
    ]
    for label, value in rows:
        ws_summary.append([label, value])

    ws_or = wb.create_sheet("OR")
    _write_header(
        ws_or,
        [
            "Cirujano",
            "Cirugías",
            "RCP",
            "Tasa RCP (%)",
            "OR",
            "IC95% low",
            "IC95% high",
            "p-value",
            "RR",
            "Diferencia absoluta de riesgo",
            "Corrección Haldane",
        ],
    )
    for row in bundle.surgeon_rows:
        ws_or.append(
            [
                row.surgeon_name,
                row.surgeries,
                row.pcr_events,
                row.pcr_rate_pct,
                row.odds_ratio if row.odds_ratio is not None else "N/A",
                row.ci_low if row.ci_low is not None else "N/A",
                row.ci_high if row.ci_high is not None else "N/A",
                row.p_value if row.p_value is not None else "N/A",
                row.relative_risk if row.relative_risk is not None else "N/A",
                row.absolute_risk_difference if row.absolute_risk_difference is not None else "N/A",
                "Sí" if row.used_haldane_correction else "No",
            ]
        )

    ws_dist = wb.create_sheet("Distribución")
    _write_header(
        ws_dist,
        ["Cirujano", "RCP", "Riesgo suavizado (%)", "Peso", "Casos sugeridos"],
    )
    for row in bundle.distribution:
        ws_dist.append(
            [
                row.surgeon_name,
                row.pcr_events,
                row.smoothed_risk_pct,
                row.weight,
                row.suggested_cases,
            ]
        )

    ws_risk = wb.create_sheet("Factores de riesgo")
    _write_header(
        ws_risk,
        ["Factor", "Casos", "Prevalencia (%)", "Porcentaje acumulado (%)"],
    )
    for row in bundle.risk_factors:
        ws_risk.append([row.label, row.cases, row.prevalence_pct, row.cumulative_pct])

    ws_reint = wb.create_sheet("Reintervenciones")
    _write_header(ws_reint, ["Tipo general", "Número"])
    for row in bundle.reintervention_types:
        ws_reint.append([row.label, row.count])
    ws_reint.append(["Total", sum(r.count for r in bundle.reintervention_types)])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
